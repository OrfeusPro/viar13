from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    r"C:\OSPanel\domains\asoft\viar\docs\upgrade-laravel13-filament5\Дорожная_карта_перехода_на_Filament_для_клиента_v1.0.docx"
)

# standard_business_brief token map
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(1.0)
HEADER_DISTANCE = Inches(0.492)
FOOTER_DISTANCE = Inches(0.492)
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP = 80
CELL_MARGIN_BOTTOM = 80
CELL_MARGIN_START = 120
CELL_MARGIN_END = 120

FONT = "Calibri"
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6B76"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GOLD = "7A5A00"
GREEN = "2F6B4F"
RED = "9B1C1C"
BORDER = "C9D2DC"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, *, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == CONTENT_WIDTH_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(
                cell,
                CELL_MARGIN_TOP,
                CELL_MARGIN_START,
                CELL_MARGIN_BOTTOM,
                CELL_MARGIN_END,
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill=CALLOUT, border_color=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border_color)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(el)
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(doc, *, num_id, abstract_id, fmt, text, bullet_font=None):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, suff, lvl_jc, p_pr])
    if bullet_font:
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), bullet_font)
        r_fonts.set(qn("w:hAnsi"), bullet_font)
        r_pr.append(r_fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def configure_section(section):
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE


def configure_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    left = p.add_run("ViarCanvas | дорожная карта миграции")
    set_run_font(left, size=9, bold=True, color=MUTED)
    p.add_run("\t")
    right = p.add_run("20 июля 2026")
    set_run_font(right, size=9, color=MUTED)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    label = fp.add_run("Страница ")
    set_run_font(label, size=9, color=MUTED)
    add_page_field(fp)


def add_body_paragraph(doc, text="", *, bold_lead=None, color=None, keep=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color)
    p.paragraph_format.keep_together = keep
    return p


def add_bullet(doc, text, num_id=40):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    apply_numbering(p, num_id)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_numbered(doc, text, num_id=41):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    apply_numbering(p, num_id)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_callout(doc, label, text, *, fill=CALLOUT, border=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = True
    set_paragraph_shading(p, fill=fill, border_color=border)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, bold=True, color=NAVY)
    r2 = p.add_run(text)
    set_run_font(r2, color=NAVY)
    return p


def fill_cell(cell, text, *, bold=False, color=NAVY, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)


def add_metadata_table(doc):
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [4680, 4680])
    set_table_borders(table, color="D8E0E8", size="4")
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    set_row_cant_split(table.rows[1])
    data = [
        ("Статус", "Аудит завершён; подготовка новой платформы начата"),
        ("Первый приоритет", "Заказы и CRM"),
        ("Принцип", "Работающий сайт не останавливается на время разработки"),
        ("Исполнитель", "Один программист; письменный отчёт раз в неделю"),
    ]
    for idx, (label, value) in enumerate(data):
        row = idx // 2
        col = idx % 2
        cell = table.cell(row, col)
        shade_cell(cell, LIGHT_GRAY if row == 0 else WHITE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        lr = p.add_run(f"{label}\n")
        set_run_font(lr, size=9, bold=True, color=BLUE)
        vr = p.add_run(value)
        set_run_font(vr, size=10.2, color=NAVY)
    return table


def add_roadmap_table(doc):
    rows = [
        ("0. Основа", "Тестовая среда, новая панель, безопасные права, мониторинг и возврат.", "160–280 ч\n4–7 нед.", "В работе"),
        ("1. Заказы и CRM", "Заказы, статусы, CRM-действия и история работают эквивалентно текущей системе.", "200–360 ч\n5–9 нед.", "Первый приоритет"),
        ("2. Пользователи и роли", "Вход, личный кабинет и видимость функций соответствуют текущим правам.", "100–180 ч\n2,5–4,5 нед.", "По очереди"),
        ("3. Чаты", "Сообщения, вложения и права участников сохраняются без потери истории.", "80–160 ч\n2–4 нед.", "По очереди"),
        ("4. Каталог и цены", "Товары, услуги, варианты, цены и связи переносятся с проверкой расчётов.", "140–260 ч\n3,5–6,5 нед.", "По очереди"),
        ("5. Переводы", "Сохраняются активные языки, ключи и согласованная production-редакция.", "80–140 ч\n2–3,5 нед.", "По очереди"),
        ("6. Платежи и доставка", "Текущее поведение провайдеров сохраняется и включается после безопасных проверок.", "120–220 ч\n3–5,5 нед.", "По очереди"),
        ("7. Контент и SEO", "Контент, блог, галерея, SEO и служебные страницы переходят в новую панель.", "160–300 ч\n4–7,5 нед.", "По очереди"),
        ("8. Финальная приёмка", "Общая проверка, стабилизация, переключение и последующий вывод Voyager.", "120–180 ч\n3–4,5 нед.", "В финале"),
        ("Итого", "Полный подтверждённый объём миграции при работе одного программиста.", "1 160–2 080 ч\n29–52 инж. нед.", "Уточняется по факту"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [1720, 4400, 1700, 1540]
    set_table_geometry(table, widths)
    set_table_borders(table)
    headers = ["Этап", "Результат для клиента", "Оценка", "Статус"]
    for idx, text in enumerate(headers):
        shade_cell(table.rows[0].cells[idx], BLUE_GRAY)
        fill_cell(table.rows[0].cells[idx], text, bold=True, color=NAVY, size=9.5)
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for ridx, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            shade_cell(cells[idx], LIGHT_GRAY if ridx == len(rows) else WHITE)
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in (2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            color = GREEN if idx == 3 and row[3] == "В работе" else NAVY
            fill_cell(cells[idx], text, bold=(ridx == len(rows) or idx == 0), color=color, size=9.2, align=align)
        set_row_cant_split(table.rows[-1])
        set_table_geometry(table, widths)
    return table


def build():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
        configure_header_footer(section)

    add_numbering_definition(doc, num_id=40, abstract_id=40, fmt="bullet", text="•", bullet_font="Symbol")
    add_numbering_definition(doc, num_id=41, abstract_id=41, fmt="decimal", text="%1.")
    add_numbering_definition(doc, num_id=42, abstract_id=42, fmt="decimal", text="%1.")

    props = doc.core_properties
    props.title = "Дорожная карта перехода с Voyager на Filament"
    props.subject = "Понятный план миграции для клиента"
    props.author = "ViarCanvas project team"
    props.keywords = "Laravel 13, Filament 5, Voyager, миграция, дорожная карта"
    props.comments = "Клиентская версия без технических приложений"

    # Customer-pack opening block (named override over standard_business_brief).
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("ДОРОЖНАЯ КАРТА ДЛЯ КЛИЕНТА")
    set_run_font(r, size=10, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Переход с Voyager на Filament")
    set_run_font(r, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Как заменить админ-панель без остановки сайта, потери заказов, прав и переводов")
    set_run_font(r, size=13.5, color=MUTED)

    add_metadata_table(doc)
    add_callout(
        doc,
        "Коротко",
        "мы не переписываем весь сайт одновременно. Новая админ-панель создаётся рядом со старой, а функции переносятся по модулям. Публичная часть сайта и текущие стили остаются без изменений. Старый модуль отключается только после проверки, приёмки и готового пути возврата.",
    )

    doc.add_heading("1. Цель проекта", level=1)
    add_body_paragraph(
        doc,
        "Voyager больше не является надёжной основой для дальнейшего обновления Laravel. Цель проекта — перейти на поддерживаемую админ-панель Filament и сохранить привычную бизнес-логику сайта.",
    )
    add_bullet(doc, "Перенести все пункты текущего административного меню, связанные кнопки, действия и специальные страницы.")
    add_bullet(doc, "Сохранить одну общую панель с тем же разграничением доступа по ролям.")
    add_bullet(doc, "Сохранить работу заказов, CRM, чатов, каталога, переводов, платежей, доставки, контента и SEO.")
    add_bullet(doc, "Получить основу, которую можно дальше обновлять вместе с Laravel и Filament.")

    doc.add_heading("2. Почему переход на Filament — это большой объём работ", level=1)
    add_body_paragraph(
        doc,
        "Filament даёт современную основу интерфейса, но не переносит автоматически настройки Voyager, собственные страницы проекта и бизнес-логику. Новая панель может выглядеть стандартно, однако её поведение нужно заново связать с действующими данными, ролями и интеграциями.",
    )
    add_callout(
        doc,
        "Главное отличие",
        "установка Filament создаёт пустой каркас. Она не создаёт готовые Заказы, CRM, чаты, переводы, платежи и остальные разделы ViarCanvas.",
        fill="FFF8E8",
        border=GOLD,
    )
    add_bullet(doc, "В текущей админке обнаружено 133 административных модуля, около 2 000 настроенных полей и 138 пунктов меню.")
    add_bullet(doc, "Помимо стандартных разделов существуют 58 собственных административных маршрутов и 40 переопределённых экранов Voyager.")
    add_bullet(doc, "Нужно сохранить восемь ролей, 675 прав и более 2 000 существующих назначений прав без случайного расширения доступа.")
    add_bullet(doc, "Для каждого раздела отдельно создаются списки, фильтры, формы, проверки данных, кнопки, массовые действия и сообщения об ошибках.")
    add_bullet(doc, "Собственные операции с заказами, файлами, переводами, платежами и внешними сервисами нельзя заменить стандартными компонентами без адаптации.")
    add_bullet(doc, "После реализации каждого раздела требуется сравнение со старой системой, проверка ролей, приёмка и подготовленный возврат.")

    doc.add_heading("3. Что останется без изменений", level=1)
    add_bullet(doc, "Публичный сайт, его дизайн, ручные стили и текущая сборка frontend не переводятся на Vite и не перегенерируются.")
    add_bullet(doc, "Production остаётся основным источником данных; новые заказы, чаты и переводы продолжают создаваться во время разработки.")
    add_bullet(doc, "Сохраняются идентификаторы заказов, статусы и правило orders.id = lead_id для CRM.")
    add_bullet(doc, "Сохраняются активные языки и переводы; при расхождениях основной считается production-редакция.")
    add_bullet(doc, "Платежи, доставка, почта и внешние интеграции не переключаются до отдельной проверки.")

    doc.add_heading("4. Что уже сделано", level=1)
    add_bullet(doc, "Проведён полный аудит кода, маршрутов, административного меню, ролей, прав, интеграций и переводов.")
    add_bullet(doc, "Составлена карта текущей админки: 133 административных модуля, 138 пунктов меню, 58 нестандартных маршрутов, 8 ролей и 675 прав.")
    add_bullet(doc, "Переводы production и локальной версии сверены; область расхождений известна и не требует массовой замены словарей.")
    add_bullet(doc, "Отдельно собран и протестирован базовый проект Laravel 13 / Filament 5; он не подключён к рабочим данным и внешним сервисам.")
    add_bullet(doc, "Проверен сценарий резервирования и возврата test-сайта; test продолжает работать в прежней конфигурации.")
    add_callout(
        doc,
        "Текущий статус",
        "подготовительный этап уже начат, но новая админка пока не включена для пользователей и не имеет права изменять production-данные.",
        fill="EEF6F1",
        border=GREEN,
    )

    doc.add_heading("5. Как будет проходить перенос", level=1)
    add_numbered(doc, "Готовим безопасную основу и отдельную тестовую среду (staging), не меняя работающий сайт.")
    add_numbered(doc, "Переносим выбранный модуль сначала в режиме чтения и сравниваем его данные со старой системой.")
    add_numbered(doc, "Проверяем роли, действия, ошибки, интеграции и пользовательские сценарии; фиксируем результат приёмки.")
    add_numbered(doc, "Включаем запись только для принятого модуля. В один момент времени владельцем записи остаётся либо Voyager, либо Filament.")
    add_numbered(doc, "Наблюдаем за модулем и при необходимости возвращаем только его на старую систему, не останавливая весь сайт.")

    doc.add_heading("6. Очерёдность и предварительная оценка", level=1)
    add_body_paragraph(
        doc,
        "Оценка дана для одного программиста. Инженерная неделя — это ориентир трудоёмкости, а не обещанная календарная дата: поддержку production, ожидание приёмки и внешние задачи нужно учитывать отдельно.",
    )
    add_roadmap_table(doc)
    add_callout(
        doc,
        "Важно",
        "этапы принимаются отдельно. После каждого этапа уточняется оценка следующего — по фактическим результатам, а не по предположениям.",
    )

    doc.add_heading("7. Что именно входит в функциональные этапы", level=1)
    add_body_paragraph(
        doc,
        "Ниже перечислен клиентский результат каждого крупного блока. Внутри каждого блока дополнительно выполняются настройка прав, обработка ошибок, журналирование, автоматические проверки и сценарий возврата.",
    )

    doc.add_heading("Заказы и CRM", level=2)
    add_bullet(doc, "Список заказов, поиск, фильтры, карточка заказа, статусы, ответственные и связанные данные клиента.")
    add_bullet(doc, "Все существующие кнопки и действия менеджера, изображения, комментарии, документы и связь с CRM.")
    add_bullet(doc, "Сохранение orders.id = lead_id, защита от повторной обработки и контроль переходов между статусами.")

    doc.add_heading("Пользователи, роли и авторизация", level=2)
    add_bullet(doc, "Вход, регистрация, Socialite, личный кабинет, восстановление доступа и существующее поведение сессий.")
    add_bullet(doc, "Карточки пользователей, восемь текущих ролей, видимость меню, прямые ссылки и разрешённые действия.")
    add_bullet(doc, "Запрет доступа по умолчанию для неизвестных или неподтверждённых действий.")

    doc.add_heading("Чаты", level=2)
    add_bullet(doc, "Переписка клиента, менеджера и художника, история сообщений, статусы и привязка к заказу.")
    add_bullet(doc, "Вложения, уведомления, проверка владельца сообщения и защита от повторной отправки.")

    doc.add_heading("Каталог и цены", level=2)
    add_bullet(doc, "Товары и услуги, варианты, размеры, страны, цены, дополнительные опции и связи между сущностями.")
    add_bullet(doc, "Медиа, порядок отображения, расчёты стоимости и проверка результатов на известных примерах.")

    doc.add_heading("Переводы", level=2)
    add_bullet(doc, "Сохранение всех активных языков, переводов из базы и файловых словарей без массовой перегенерации.")
    add_bullet(doc, "Версионная публикация, проверка ключевых страниц на каждом языке и атомарный возврат предыдущего пакета.")
    add_bullet(doc, "Production-редакция остаётся основной; dormant CN/JP-файлы архивируются и не включаются без отдельного решения.")

    doc.add_heading("Платежи, доставка и внешние сервисы", level=2)
    add_bullet(doc, "PayPal, Paysera/WebToPay, Venipak, SMTP, Synvolve/SA, Facebook и Google сохраняют текущее поведение.")
    add_bullet(doc, "Проверяются запросы, ответы, повторные callback, ошибки провайдера и неизвестный результат операции.")
    add_bullet(doc, "Реальная запись переключается только после доступных sandbox-проверок и контролируемого smoke-теста.")

    doc.add_heading("Контент, блог, галерея и SEO", level=2)
    add_bullet(doc, "Редактирование страниц, публикаций, изображений, галерей, отзывов и вспомогательного контента.")
    add_bullet(doc, "SEO-поля, URL, redirects, sitemap, feeds и проверки, исключающие новые 404 или циклы перенаправлений.")

    doc.add_heading("8. Первый практический результат", level=1)
    add_body_paragraph(
        doc,
        "Первый приоритет — Заказы и CRM. До 1 октября 2026 года реалистичная базовая цель — завершить фундамент и показать на staging рабочий контур Заказов и CRM как минимум в безопасном режиме чтения, с проверенными ролями и журналом действий.",
    )
    add_bullet(doc, "Ограниченное включение части Заказов и CRM в production возможно только после прохождения всех проверок и при отсутствии блокирующих дефектов.")
    add_bullet(doc, "Полный отказ от Voyager к 1 октября не обещается: остальные модули переходят последовательно после первой поставки.")
    add_callout(
        doc,
        "Ориентир, а не гарантия",
        "дата зависит от фактической сложности legacy-логики, поддержки работающего сайта и скорости приёмки. Прогноз обновляется раз в неделю.",
        fill="FFF8E8",
        border=GOLD,
    )

    doc.add_heading("9. Когда этап считается завершённым", level=1)
    add_bullet(doc, "Все согласованные функции модуля доступны и ведут себя так же, как в текущей системе.")
    add_bullet(doc, "Роли видят только разрешённые разделы, записи и действия; новые права автоматически не выдаются.")
    add_bullet(doc, "Данные, статусы, переводы и файлы сверены; необъяснимых потерь или дублей нет.")
    add_bullet(doc, "Интеграции модуля проверены безопасным способом, а ошибки можно отследить.")
    add_bullet(doc, "Есть проверенный способ возврата на старый модуль.")
    add_bullet(doc, "Программист и ответственный за приёмку письменно фиксируют PASS.")

    doc.add_heading("10. Основные риски и как они контролируются", level=1)
    add_body_paragraph(doc, "Работающий production. ", bold_lead="Работающий production. ")
    add_body_paragraph(doc, "Во время разработки появляются новые заказы и переводы. Поэтому production остаётся источником истины, а перед каждым переключением выполняется финальная сверка изменений.")
    add_body_paragraph(doc, "Скрытая legacy-логика. ", bold_lead="Скрытая legacy-логика. ")
    add_body_paragraph(doc, "Админка содержит собственные страницы и действия, которых нет в стандартном Voyager. Перенос идёт по функциям и маршрутам, с тестами каждого сценария.")
    add_body_paragraph(doc, "Права доступа. ", bold_lead="Права доступа. ")
    add_body_paragraph(doc, "Новая панель закрывает неизвестные действия по умолчанию. Восемь существующих ролей сравниваются со старой системой до включения.")
    add_body_paragraph(doc, "Переводы. ", bold_lead="Переводы. ")
    add_body_paragraph(doc, "Массовая перегенерация запрещена. Сохраняются все активные языки, а изменения публикуются версионным пакетом с возможностью возврата.")
    add_body_paragraph(doc, "Платежи и доставка. ", bold_lead="Платежи и доставка. ")
    add_body_paragraph(doc, "Не все провайдеры имеют полноценный sandbox. До контролируемой проверки владельцем записи остаётся старая система; после включения действует усиленный мониторинг.")
    add_body_paragraph(doc, "Нулевая плановая пауза. ", bold_lead="Нулевая плановая пауза. ")
    add_body_paragraph(doc, "Нельзя выполнять длительную общую миграцию. Используется модульное переключение в субботу вечером и возврат отдельного модуля при проблеме.")

    doc.add_heading("11. Что требуется от клиента", level=1)
    add_bullet(doc, "Утверждать состав и оценку каждого этапа до начала реализации.")
    add_bullet(doc, "Участвовать в приёмке бизнес-сценариев либо назначить ответственного; техническую приёмку и переводы ведёт программист.")
    add_bullet(doc, "Быть доступным для связи во время согласованного переключения; техническое переключение выполняет программист.")
    add_bullet(doc, "Согласовывать новые пожелания отдельно, если они меняют текущее поведение или объём этапа.")
    add_bullet(doc, "Принимать еженедельный отчёт: что завершено, сколько часов затрачено, какие риски обнаружены и что планируется дальше.")

    doc.add_heading("12. Что не входит в текущую миграцию", level=1)
    add_bullet(doc, "Редизайн публичного сайта и перевод существующего frontend на Vite.")
    add_bullet(doc, "Автоматическое добавление новых функций, которых нет в текущей админке.")
    add_bullet(doc, "Одновременная замена всех внешних сервисов, файлового хранения или структуры production-данных.")
    add_bullet(doc, "Фиксированная стоимость всего проекта: бюджет согласуется по этапам на основании часов.")

    doc.add_heading("13. Следующие шаги", level=1)
    add_numbered(doc, "Разместить уже собранный Laravel 13 / Filament 5 проект в закрытом серверном каталоге, не переключая работающий test.", num_id=42)
    add_numbered(doc, "Создать отдельную обезличенную базу и безопасную конфигурацию без production-ключей и реальных внешних отправок.", num_id=42)
    add_numbered(doc, "Завершить подготовительный этап: роли, сессии, переводы, чтение заказов, чатов и файлов, мониторинг и возврат.", num_id=42)
    add_numbered(doc, "Показать клиенту первый рабочий контур Заказов и CRM на staging и зафиксировать замечания.", num_id=42)
    add_numbered(doc, "После письменной приёмки отдельно согласовать первое ограниченное production-включение.", num_id=42)

    add_callout(
        doc,
        "Предлагаемое решение",
        "утвердить эту дорожную карту как рабочий план, продолжить подготовительный этап и сохранить еженедельную отчётность. Любое изменение объёма оформлять отдельно до начала работ.",
        fill="EEF6F1",
        border=GREEN,
    )

    doc.add_heading("14. Статус документа", level=1)
    add_body_paragraph(
        doc,
        "Клиентская версия 1.0 от 20 июля 2026 года. Документ упрощает технический план v2.36 и последующие результаты аудита. При расхождении технических деталей источником истины остаются технический аудит, реестр решений и письменная приёмка этапа.",
        color=MUTED,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
